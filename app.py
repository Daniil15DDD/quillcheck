"""QuillCheck UI. Run with ./run.sh; serves on 127.0.0.1 only."""
import os

os.environ["HF_HUB_OFFLINE"] = "1"  # detector.py enforces this too
os.environ["TRANSFORMERS_OFFLINE"] = "1"

import base64
import html
import io
import time
from collections import Counter
from pathlib import Path

import streamlit as st
from docx import Document

import calibration
import detector
import likelihood
import rewriter
import worker

HERE = Path(__file__).parent
DISCLAIMER = (
    "<b>Model estimates only — not evidence of who wrote a text.</b> "
    "Everything runs on this Mac. Text is not saved or sent anywhere. "
    "Scores are not comparable with GPTZero or any other detector."
)
LOCK = '<i class="qc-lock"></i>'  # icon drawn by ui.css (st.html strips inline SVG)

# Try-it texts. "Human-written" is public-domain text by a known author; the AI text was written
# by an AI model for this app. "Mixed" joins the two.
SMITH_1776 = (
    "The greatest improvement in the productive powers of labour, and the greatest part of the skill, dexterity, and judgment "
    "with which it is any where directed, or applied, seem to have been the effects of the division of labour. The effects of "
    "the division of labour, in the general business of society, will be more easily understood, by considering in what manner "
    "it operates in some particular manufactures. It is commonly supposed to be carried furthest in some very trifling ones; not "
    "perhaps that it really is carried further in them than in others of more importance: but in those trifling manufactures "
    "which are destined to supply the small wants of but a small number of people, the whole number of workmen must necessarily "
    "be small; and those employed in every different branch of the work can often be collected into the same workhouse, and "
    "placed at once under the view of the spectator."
)
AI_SAMPLE = (
    "In today's rapidly evolving business landscape, organisations must leverage innovative strategies to remain competitive. "
    "By fostering a culture of continuous improvement, companies can unlock new opportunities for sustainable growth and "
    "long-term value creation. Furthermore, effective stakeholder engagement plays a pivotal role in ensuring that strategic "
    "initiatives align with broader organisational objectives.\n\n"
    "Moreover, embracing digital transformation enables businesses to streamline operations, enhance customer experiences, and "
    "drive meaningful value across the organisation. By harnessing data-driven insights, leaders can make informed decisions "
    "that strengthen resilience in an increasingly complex market. Ultimately, a holistic approach that balances innovation with "
    "operational excellence positions organisations to thrive in a dynamic and interconnected global economy."
)
SAMPLES = {
    "AI-written": AI_SAMPLE,
    "Human-written (Adam Smith, 1776)": SMITH_1776,
    "Mixed": SMITH_1776 + "\n\n" + AI_SAMPLE.split("\n\n")[1],
}

st.set_page_config(page_title="QuillCheck", page_icon=str(HERE / "desktop" / "icon.png"), layout="wide")
dark = st.context.theme.type == "dark"
css = (HERE / "ui.css").read_text(encoding="utf-8")
st.html(f"<style>{css}</style>" + ('<div class="qc-dark"></div>' if dark else ""))  # ui.css switches palette on .qc-dark
logo = base64.b64encode((HERE / "docs" / "assets" / "logo.svg").read_bytes()).decode()
st.html(
    f'<div class="qc-head"><img class="qc-logo" alt="" src="data:image/svg+xml;base64,{logo}">'
    '<div><div class="name">QuillCheck</div><div class="what">Check and edit your writing, privately on this Mac.</div></div>'
    f'<span class="qc-badge">{LOCK} Runs offline</span></div>'
    f'<p class="qc-fine">{DISCLAIMER}</p>'
)


def use_sample(key: str) -> None:
    choice = st.session_state.get(f"{key}_sample")
    if choice:
        st.session_state[f"{key}_text"] = SAMPLES[choice]
        st.session_state[f"{key}_sample"] = None


def read_input(key: str, samples: bool = False) -> str:
    """The compose card's text box and file upload; returns the text to use."""
    with st.container(horizontal=True, vertical_alignment="center", key=f"{key}_top"):
        st.html('<span class="qc-hint">Paste your text or</span>')
        upload = st.file_uploader("Upload a .docx or .txt file", type=["docx", "txt"], key=f"{key}_file", label_visibility="collapsed")
    text = st.text_area("Paste text", height=260, key=f"{key}_text", label_visibility="collapsed",
                        placeholder="Paste an essay, report or paragraph here…")
    if samples and not upload:
        st.pills("Try a sample", list(SAMPLES), key=f"{key}_sample", on_change=use_sample, args=(key,))
    if upload is not None:
        try:
            if upload.name.lower().endswith(".docx"):
                doc = Document(io.BytesIO(upload.getvalue()))
                text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
            else:
                text = upload.getvalue().decode("utf-8", errors="replace")
        except Exception as e:  # corrupt or mislabelled file
            st.error(f"Could not read {upload.name}: {type(e).__name__}. Paste the text instead.")
            return ""
        st.caption(
            f"Using {len(text.split())} words from **{upload.name}** instead of the text box. "
            "From .docx only body paragraphs are read; tables, text boxes, headers, footers and footnotes are skipped."
        )
    return text.strip()


def count_line(text: str) -> None:
    words = len(text.split())
    st.html(f'<span class="qc-count">{words:,} words · {len(text):,} characters'
            f'{" · 150+ words gives steadier results" if 0 < words < 150 else ""}</span>')


def scroll_to(selector: str) -> None:
    """Scroll the page to fresh results, so a finished run doesn't look like it's still working."""
    # Waits up to 3 s for the results, then lets the layout settle; a smooth scroll gets cancelled while it grows.
    # The run id keeps the frame unique, otherwise Streamlit reuses it and the script doesn't run again.
    st.iframe(f"<script>/* {time.time_ns()} */ let n = 0; const t = setInterval(() => {{ const el = window.parent.document.querySelector({selector!r});"
              f" if (el || ++n > 30) {{ clearInterval(t); setTimeout(() => el?.scrollIntoView({{block: 'start'}}), 600); }} }}, 100)</script>",
              height=1)  # st.iframe rejects 0


def run(text: str, detailed: bool) -> detector.Result | None:
    return run_in_worker("Analysing locally…", lambda p: detector.analyse_in_subprocess(text, detailed=detailed, progress=p))


def run_in_worker(label: str, call):
    error = None
    slot = st.empty()
    with slot.container(), st.status(label, expanded=False) as status:
        try:
            result = call(lambda msg: status.update(label=msg))
        except worker.WorkerError as e:
            status.update(label="Failed", state="error")
            error, result = e, None
    if error is None:
        slot.empty()  # the result itself shows it finished
    # Shown outside the collapsed status box so the explanation is always visible.
    if error is not None:
        if error.kind == "models-missing":
            st.error(f"Model files could not be loaded from the local cache ({error}). The app is offline by design; re-run the one-time model download, then try again.")
        elif error.kind == "crashed":
            st.error(f"{error} This usually means macOS ran out of memory. Close other apps or check a shorter text.")
        else:
            st.error(f"The model run failed ({error}). If this mentions memory, close other apps or check a shorter text.")
    return result


def sentence_view(result: detector.Result) -> str:
    """Highlight sentences in AI-looking paragraphs (see likelihood.flag_sentences)."""
    scored = sorted(s.binoculars for s in result.sentences if s.binoculars is not None)
    flags = likelihood.flag_sentences(result)
    rank = {v: i + 1 for i, v in reversed(list(enumerate(scored)))}
    parts, last_paragraph = ["<p>"], 0
    for s in result.sentences:
        paragraph = next(p.index for p in result.paragraphs if p.start <= s.start < p.end)
        if paragraph != last_paragraph:
            parts.append("</p><p>")
            last_paragraph = paragraph
        body = html.escape(s.text)
        if s.binoculars is None:
            parts.append(f'<span class="qc-faded" title="too short to score">{body}</span> ')
        else:
            tip = f"Relative rank {rank[s.binoculars]} of {len(scored)} in this text (1 = most AI-like here). Raw Binoculars {s.binoculars:.3f}. Not a probability."
            if s.index in flags:
                parts.append(f'<mark title="{tip}">{body}</mark> ')
            else:
                parts.append(f'<span title="{tip}">{body}</span> ')
    return '<div class="qc-page">' + "".join(parts) + "</p></div>"


BANDS = [("success", "Leans human", "0–29%"), ("warning", "Unclear", "30–69%"), ("error", "Leans AI", "70–100%")]


def score_panel(pct: int, label: str, severity: str, marked: int, total: int) -> str:
    """Ring gauge, verdict and the three bands of likelihood.reading()."""
    chips = "".join(
        f'<span class="qc-chip {sev}{" on" if sev == severity else ""}"><b>{name}</b> {rng}</span>' for sev, name, rng in BANDS
    )
    marked_line = (f"{marked} of {total} sentences highlighted for a reread" if marked
                   else "No sentences highlighted")
    return (
        '<div class="qc-score">'
        f'<div class="qc-ring {severity}" style="--pct:{pct}"><div class="qc-ring-text"><b>{pct}%</b><span>AI likelihood</span></div></div>'
        f'<div class="qc-verdict"><div class="qc-verdict-label">Overall reading</div>'
        f'<div class="qc-verdict-text {severity}">{html.escape(label)}</div>'
        f'<div class="qc-verdict-sub">{marked_line}</div></div></div>'
        f'<div class="qc-chips">{chips}</div>'
        f'<p class="qc-caveat">{html.escape(likelihood.CAVEAT)}</p>'
    )


def top_sentences(result: detector.Result) -> str:
    flags = likelihood.flag_sentences(result)
    marked = sorted((s for s in result.sentences if s.index in flags), key=lambda s: s.binoculars)
    if not marked:
        return '<p class="qc-empty">Nothing stands out: no paragraph looked AI-generated, so no sentences are marked.</p>'
    rows = "".join(
        f'<div class="qc-row"><span class="qc-dots" title="Rank {i + 1} of {len(marked)} marked">{"<i></i>" * max(1, 3 - (3 * i) // len(marked))}</span>'
        f'<span>{html.escape(s.text)}</span></div>' for i, s in enumerate(marked)
    )
    return f'<div class="qc-list">{rows}</div>'


REWRITE_NOTICE = (
    "<b>AI-assisted editing — review every change before you use it.</b> A small local model makes these suggestions; "
    "it can still change meaning or make mistakes. It is not designed, tuned or tested to change AI-detector results "
    "and makes no claim about GPTZero or any other detector. If you use it for assessed work, follow your university's "
    "rules on AI use and acknowledge it where required."
)
STATUS_LABELS = {
    "edited": "Edited",
    "review": "Edited — read the notes",
    "kept-original": "Kept your original (a check failed) — suggestion shown for reference only",
    "unchanged": "No changes suggested",
}


def rewrite_view(res: rewriter.RewriteResult) -> None:
    counts = Counter(c.status for c in res.chunks)
    stats = [("Edited", counts["edited"]), ("Read the notes", counts["review"]),
             ("Kept original", counts["kept-original"]), ("No changes", counts["unchanged"])]
    st.html('<div class="qc-stats">' + "".join(f'<div><b>{n}</b><span>{name}</span></div>' for name, n in stats) + "</div>")
    st.caption(
        f"**What was processed:** all {res.paragraphs} paragraph(s), {res.words_in} words, in {len(res.chunks)} chunk(s) of up to "
        f"{rewriter.CHUNK_TOKENS} tokens (long paragraphs are split between sentences). Nothing was truncated. "
        f"Style: {rewriter.MODES[res.mode][0]}."
    )
    if counts["kept-original"]:
        st.warning(
            f"{counts['kept-original']} chunk(s) keep your original wording because a check failed "
            "(a number, citation or link was lost or added, or the output was cut off). The model's version is shown for reference only."
        )
    rows = []
    for c in res.chunks:
        left, right = rewriter.diff_html(c.original, c.suggested or c.original)
        notes = "".join(f"<li>{html.escape(msg)}</li>" for _, msg in c.issues)
        rows.append(
            f"<tr><td>{left}</td><td>{right}<div class='rw-status rw-{c.status}'>{STATUS_LABELS[c.status]}</div>"
            f"{f'<ul class=rw-notes>{notes}</ul>' if notes else ''}</td></tr>"
        )
    st.subheader("Side by side")
    st.caption("Left: your text, with removed or replaced words struck through. Right: the suggestion, with new words highlighted.")
    st.markdown("<table class='rw'><tr><th>Original</th><th>Suggestion</th></tr>" + "".join(rows) + "</table>", unsafe_allow_html=True)
    st.subheader("Rewritten text — AI-assisted, review before use")
    st.caption("Copy with the icon at the top right of the box, or download it. Chunks that failed a check keep your original wording here.")
    st.code(res.output, language=None, wrap_lines=True)
    st.download_button("Download as .txt", res.output, file_name="rewritten.txt", mime="text/plain", on_click="ignore")


check, rewrite_tab, calibrate, about = st.tabs(["Check", "Rewrite", "Calibrate", "About"])

with rewrite_tab:
    st.html(f'<div class="qc-note">{REWRITE_NOTICE}</div>')
    with st.container(border=True, key="rewrite_card"):
        mode = st.radio("Editing style", list(rewriter.MODES), format_func=lambda k: rewriter.MODES[k][0], horizontal=True, key="rewrite_mode")
        st.caption(rewriter.MODES[mode][1] + " Australian English. Facts, numbers, names and citations must stay unchanged.")
        source = read_input("rewrite")
        with st.container(horizontal=True, vertical_alignment="center", horizontal_alignment="distribute", key="rewrite_foot"):
            count_line(source)
            go = st.button("Rewrite", type="primary", disabled=not source)
    if go:
        res = run_in_worker("Editing locally…", lambda p: rewriter.rewrite_in_subprocess(source, mode, progress=p))
        if res is not None:
            st.session_state["rewrite_result"] = (source, mode, res)
    stored = st.session_state.get("rewrite_result")
    if stored and stored[0] == source and stored[1] == mode:
        rewrite_view(stored[2])
        if go:
            scroll_to(".qc-stats")

with check:
    with st.container(border=True, key="check_card"):
        text = read_input("check", samples=True)
        with st.container(horizontal=True, vertical_alignment="center", horizontal_alignment="distribute", key="check_foot"):
            count_line(text)
            go = st.button("Analyse", type="primary", disabled=not text)
    st.html(f'<p class="qc-private">{LOCK} Your text stays on this Mac. Nothing is uploaded or saved.</p>')

    if go:
        result = run(text, detailed=True)
        if result is not None:
            st.session_state["check_result"] = (text, result)
    stored = st.session_state.get("check_result")
    if stored and stored[0] == text:
        result = stored[1]
        for warning in result.warnings:
            st.warning(warning)
        flags = likelihood.flag_sentences(result)
        if go:
            scroll_to(".st-key-check_text_card")
        left, right = st.columns([1.35, 1], gap="large")
        with left, st.container(border=True, key="check_text_card"):
            st.html('<div class="qc-card-head">Your text'
                    f'{" <span class=qc-legend><mark>Highlighted</mark> = worth a reread, not a verdict</span>" if flags else ""}</div>')
            st.markdown(sentence_view(result), unsafe_allow_html=True)
        with right, st.container(border=True, key="check_result_card"):
            reading = likelihood.reading(result.classifier)
            if reading is None:
                st.info("No estimate: the text was too short to score.")
            else:
                scored = sum(s.binoculars is not None for s in result.sentences)
                st.html(score_panel(*reading, len(flags), scored))
                human_clf = calibration.scores("human", "classifier")
                if len(human_clf) >= 3:
                    st.caption(f"For comparison, your own writing samples scored {human_clf[0] * 100:.0f}%–{human_clf[-1] * 100:.0f}%.")

            sent_tab, para_tab, detail_tab = st.tabs(["Sentences", "Paragraphs", "Details"])
            with sent_tab:
                st.html(top_sentences(result))
                st.caption(likelihood.HIGHLIGHT_NOTE + " Single-sentence scores are noisy, so a mark is a reason to reread, **not** a verdict. "
                           "Hover a sentence in your text for its rank; faded text was too short to score.")
            with para_tab:
                st.dataframe(
                    [
                        {
                            "#": p.index + 1,
                            "Starts with": " ".join(p.text.split()[:6]) + "…",
                            "AI likelihood %": None if p.classifier is None else round(p.classifier * 100),
                        }
                        for p in result.paragraphs
                    ],
                    hide_index=True,
                    width="stretch",
                    column_config={
                        "#": st.column_config.NumberColumn(width="small"),
                        "AI likelihood %": st.column_config.ProgressColumn(min_value=0, max_value=100, format="%d%%"),
                    },
                )
            with detail_tab:
                st.metric("Classifier score (desklib)", "—" if result.classifier is None else f"{result.classifier * 100:.0f}%")
                st.caption("The number shown above. A model trained on news, books, abstracts, reviews, Reddit, recipes, Wikipedia and poetry, not coursework. 50% is the model card's cut-off.")
                st.metric("Binoculars raw score — EXPERIMENTAL (lower = more AI-like)", "—" if result.binoculars is None else f"{result.binoculars:.3f}")
                st.caption("A second, independent check. It has no validated threshold for these models, so it is not used for the headline and only means something next to your own samples.")
                context = calibration.binoculars_context(result.binoculars) if result.binoculars is not None else None
                if context is None:
                    st.caption("Uncalibrated: add at least 3 samples of your own writing in **Calibrate** to compare against.")
                else:
                    getattr(st, context[0])(context[1])
                cov = result.coverage
                if cov:
                    st.caption(
                        f"**What was scored:** Binoculars read all {cov.get('binoculars_tokens', 0)} tokens in {cov.get('binoculars_windows', 0)} window(s) of up to "
                        f"{detector.WINDOW_TOKENS} tokens (context does not carry across windows). The classifier read {cov.get('classifier_tokens', 0)} of "
                        f"{cov.get('classifier_total', 0)} tokens in {cov.get('classifier_windows', 0)} window(s) of up to {detector.CLASSIFIER_MAX_TOKENS} tokens "
                        "(only whitespace between paragraphs can be skipped). No text was truncated."
                    )
    elif not text:
        st.html(
            '<div class="qc-how">'
            '<div><b>1</b><span>Paste text or upload a .docx</span></div>'
            '<div><b>2</b><span>Press Analyse. The first run loads the models (about 20 seconds)</span></div>'
            '<div><b>3</b><span>Read the overall estimate and the highlighted sentences</span></div></div>'
        )

with calibrate:
    st.html(
        '<div class="qc-note">Score texts whose origin you know, so the Binoculars number has a personal reference point. '
        "Best: 5–10 of your own assignments written <b>before you used AI tools</b>. Optionally add a few AI-written samples you generate yourself. "
        "Only the scores are saved (in <code>calibration.json</code>), never the text.</div>"
    )
    with st.container(border=True, key="calibrate_card"):
        label = st.radio("This text is…", ["human", "ai"], format_func={"human": "My own writing", "ai": "AI-generated (made by me for comparison)"}.get, horizontal=True)
        sample = read_input("calibrate")
        with st.container(horizontal=True, vertical_alignment="center", horizontal_alignment="distribute", key="calibrate_foot"):
            count_line(sample)
            add = st.button("Score and add", type="primary", disabled=not sample)
    if add:
        result = run(sample, detailed=False)
        if result is not None:
            calibration.add(label, result.binoculars, result.classifier, result.words)
            st.success("Added.")
            for warning in result.warnings:
                st.warning(warning)

    samples = calibration.load()
    if samples:
        st.subheader(f"Your samples ({len(samples)})")
        st.dataframe(samples, hide_index=True, width="stretch")
        if st.button("Delete all calibration samples"):
            calibration.save([])
            st.rerun()

with about:
    st.markdown(
        f"""
**What it runs** (offline after the first download):

- **Binoculars** (Hans et al., 2024; code BSD-3) with `{detector.OBSERVER_ID}` + `{detector.PERFORMER_ID}` (Apache-2.0).
  It compares how surprising the text is to one model against how surprising one model expects the other to find it.
  The published threshold was tuned for Falcon-7B, so with these smaller models the score only means something against your own calibration samples.
  **Experimental:** this Qwen pair has not been independently validated for Binoculars. The paper's code scores only the first 512 tokens;
  this app scores the whole text in 512-token windows and combines them, which is an extension the paper did not test.
- **`{detector.CLASSIFIER_ID}`** (MIT): a DeBERTa-v3-large classifier trained on the RAID benchmark.

**Rewrite tab:** uses the same downloaded `{rewriter.MODEL_ID}` (Apache-2.0) to suggest edits, offline. Each chunk is checked:
if a number, citation or link is lost or added, or the output is cut off, your original wording is kept. The checks cannot catch
every change in meaning (for example "creates a risk" becoming "could create a risk"), so read every change. The Rewrite tab
never runs the detector on its output and is not designed to change detector results.

**Known limits (detector):**

- Formulaic writing (reports, SWOT and other frameworks, executive summaries, reference-heavy text) looks more "AI-like" to these methods, which raises false positives.
- Perplexity-style detectors flag non-native English writers more often (Liang et al., 2023, *Patterns*).
- Short texts are unreliable. Edited, paraphrased or mixed human/AI text is hard for every detector.
- The two models can disagree. Treat disagreement as "uncertain", not as a tie-break.
- None of this is equivalent to GPTZero, Turnitin or any other product, and none of it can prove authorship.

Device: `{detector.DEVICE}`.
"""
    )
